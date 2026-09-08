#!/usr/bin/env python3
"""Generate the weekly progress deck as .pptx, with diagrams.

  python3 make_week_slides.py
Writes asap_week_20260809.pptx next to this script.

Palette and geometry match reports/make_slides.py so the decks read as one series.
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

HERE = os.path.dirname(os.path.abspath(__file__))

INK   = RGBColor(0x1A, 0x1D, 0x24)
MUTED = RGBColor(0x5B, 0x63, 0x72)
ACC   = RGBColor(0x1F, 0x6F, 0x8B)
GOOD  = RGBColor(0x1B, 0x7F, 0x53)
BAD   = RGBColor(0xA6, 0x3D, 0x40)
PALE  = RGBColor(0xE9, 0xEE, 0xF1)
PALEB = RGBColor(0xF6, 0xE7, 0xE7)
PALEG = RGBColor(0xE4, 0xF1, 0xEA)
GREY  = RGBColor(0xC9, 0xD0, 0xD6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ primitives
def tb(slide, l, t, w, h):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def run(p, text, size=18, bold=False, color=INK, after=6, align=None):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    p.space_after = Pt(after)
    if align: p.alignment = align
    return r


def header(slide, kicker, title):
    tf = tb(slide, 0.7, 0.45, 12.0, 0.35)
    run(tf.paragraphs[0], kicker.upper(), 12, True, ACC)
    tf2 = tb(slide, 0.7, 0.83, 12.0, 0.85)
    run(tf2.paragraphs[0], title, 29, True, INK)


def bullets(slide, items, top=2.0, left=0.7, width=12.0, size=16, gap=7, height=4.6):
    tf = tb(slide, left, top, width, height)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        text, color, bold = it if isinstance(it, tuple) else (it, INK, False)
        run(p, text, size, bold, color, after=gap)
    return tf


def note(slide, text, top=6.45, color=MUTED, size=13.5):
    tf = tb(slide, 0.7, top, 12.0, 0.7)
    run(tf.paragraphs[0], text, size, False, color)


def box(slide, l, t, w, h, text, fill=PALE, line=GREY, color=INK, size=13,
        bold=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(45720)
    tf.margin_top = tf.margin_bottom = Emu(18288)
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run(p, ln, size if i == 0 else size - 1.5, bold if i == 0 else False,
            color if i == 0 else MUTED, after=1, align=PP_ALIGN.CENTER)
    return sh


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    c.line.fill.solid(); c.line.fill.fore_color.rgb = color
    el = c.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    from lxml import etree
    tail = etree.SubElement(el, qn('a:tailEnd'))
    tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')
    return c


def caption(slide, l, t, w, text, color=MUTED, size=12, bold=False, align=PP_ALIGN.CENTER):
    tf = tb(slide, l, t, w, 0.42)
    run(tf.paragraphs[0], text, size, bold, color, after=0, align=align)


def table(slide, rows, top=2.2, left=0.7, col_w=None, size=14):
    ncol = len(rows[0])
    col_w = col_w or [Inches(12.0 / ncol)] * ncol
    shape = slide.shapes.add_table(len(rows), ncol, Inches(left), Inches(top),
                                   sum(col_w, Inches(0)), Inches(0.38 * len(rows)))
    t = shape.table
    for j, w in enumerate(col_w):
        t.columns[j].width = w
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            text, color, bold = cell if isinstance(cell, tuple) else (cell, INK, False)
            c = t.cell(i, j); c.text = ""
            p = c.text_frame.paragraphs[0]
            run(p, str(text), size, bold or i == 0, MUTED if i == 0 else color, 0)
            if j > 0: p.alignment = PP_ALIGN.RIGHT
    return t


# ================================================================== 1. title
s = prs.slides.add_slide(BLANK)
tf = tb(s, 0.9, 2.0, 11.5, 0.4)
run(tf.paragraphs[0], "WEEKLY PROGRESS  ·  9 AUGUST 2026", 13, True, ACC)
tf = tb(s, 0.9, 2.5, 11.6, 1.8)
run(tf.paragraphs[0], "Delta action models on 473 real dab repetitions", 40, True, INK)
tf = tb(s, 0.9, 4.15, 11.2, 1.4)
run(tf.paragraphs[0],
    "Fixed the capture pipeline, measured where the sim-to-real gap actually is, "
    "and found that ASAP's 4-DoF choice is right, for a reason the paper does not give.",
    18, False, MUTED)

# ================================================================== 2. TL;DR
s = prs.slides.add_slide(BLANK)
header(s, "summary", "Four things happened this week")
for i, (n, t1, t2, col) in enumerate([
    ("1", "Fixed three bugs in the capture data",
     "29 mm pose error  ·  1.8% timebase drift  ·  83 mm height offset", GOOD),
    ("2", "Measured where the sim-to-real gap is",
     "All four ankle joints, 4.79x worse than everything else", GOOD),
    ("3", "4-DoF delta model beats the full 23-DoF one",
     "The full model cheats. Replicated across two training seeds.", GOOD),
    ("4", "Fine-tuned a policy, and it cannot balance",
     "Two causes found. Neither is fixed yet.", BAD),
]):
    y = 1.95 + i * 1.15
    box(s, 0.7, y, 0.55, 0.75, n, fill=WHITE, line=col, color=col, size=17, bold=True)
    tf = tb(s, 1.45, y - 0.03, 11.2, 0.45)
    run(tf.paragraphs[0], t1, 18, True, INK, after=1)
    tf2 = tb(s, 1.45, y + 0.36, 11.2, 0.4)
    run(tf2.paragraphs[0], t2, 14, False, MUTED)
note(s, "ASAP publishes no working code for steps 2, 3 or 4 of its own pipeline. All three were written this week.")

# ================================================================== 3. data bugs
s = prs.slides.add_slide(BLANK)
header(s, "result 0  \u00b7  the data", "The capture data had three bugs in it")

# --- the three bugs
rows = [
    ("1", "Pose solver mis-centred on partial frames",
     "29 mm position jump whenever a marker dropped", BAD),
    ("2", "Frame rate stored as a float, then truncated",
     "48.88 became 48. A 1.8% timebase drift.", BAD),
    ("3", "Marker cluster used as the pelvis origin",
     "Every root height 83 mm too low", BAD),
]
for i, (n, t1, t2, col) in enumerate(rows):
    y = 1.95 + i * 1.25
    box(s, 0.7, y, 0.55, 0.72, n, fill=WHITE, line=col, color=col, size=16, bold=True)
    tf = tb(s, 1.45, y - 0.06, 11.3, 0.45)
    run(tf.paragraphs[0], t1, 18, True, INK, after=1)
    tf2 = tb(s, 1.45, y + 0.36, 11.3, 0.4)
    run(tf2.paragraphs[0], t2, 14, False, MUTED)
box(s, 0.7, 5.7, 12.0, 0.7,
    "All three fixed, verified, and pushed",
    fill=PALEG, line=GOOD, color=GOOD, size=15, bold=True)

note(s, "Bug 1 is explained over the next two slides. Bug 2 is the dangerous one: a wrong timebase is "
        "invisible in the training loss, and only shows up when the data is round-tripped back.",
     top=6.6, size=13.5)

# ================================================================== 3b. how the solver works
DIA2 = RGBColor(0xB4, 0x53, 0x1F)


def dot(slide, cx, cy, line, r=0.13, fill=WHITE):
    return box(slide, cx - r, cy - r, 2 * r, 2 * r, "", fill=fill, line=line,
               shape=MSO_SHAPE.OVAL)


def cross(slide, cx, cy, color, r=0.09):
    return box(slide, cx - r, cy - r, 2 * r, 2 * r, "", fill=color, line=color,
               shape=MSO_SHAPE.OVAL)


SQ = [(-0.275, -0.20), (0.275, -0.20), (-0.275, 0.20), (0.275, 0.20)]
ROT = [(-0.138, -0.273), (0.338, 0.0), (-0.338, 0.0), (0.138, 0.273)]

s = prs.slides.add_slide(BLANK)
header(s, "bug 1 explained  \u00b7  part 1", "How the pose solver works")
caption(s, 0.7, 1.7, 12.0,
        "Four markers sit on the pelvis. The solver works out how the pelvis has turned, "
        "by comparing the shape it sees now against the shape it recorded at the start.",
        MUTED, 14, False, PP_ALIGN.LEFT)

panels = [
    (0.9, "Two shapes, different place and angle",
     "Blue is the shape recorded at the start.\nOrange is what the cameras see now."),
    (5.0, "Step 1:  slide them onto the same centre",
     "This removes WHERE the robot is.\nIt does not measure any turning."),
    (9.1, "Step 2:  search for the turn that fits",
     "Try every rotation, keep the one that\nlines the markers up. That is the answer."),
]
for i, (px, title, sub) in enumerate(panels):
    caption(s, px - 0.2, 2.3, 3.9, title, INK, 13.5, True)
    cy = 3.65
    if i == 0:
        ca, cb = px + 0.75, px + 2.55
        for dx, dy in SQ: dot(s, ca + dx, cy + dy, ACC)
        cross(s, ca, cy, ACC)
        for dx, dy in ROT: dot(s, cb + dx, cy + dy, DIA2)
        cross(s, cb, cy, DIA2)
        caption(s, ca - 0.85, cy + 0.55, 1.7, "recorded", ACC, 11, True)
        caption(s, cb - 0.85, cy + 0.55, 1.7, "seen now", DIA2, 11, True)
    else:
        c = px + 1.65
        big = 0.17 if i == 2 else 0.13
        for dx, dy in SQ: dot(s, c + dx, cy + dy, ACC, r=big)
        src = ROT if i == 1 else SQ
        for dx, dy in src: dot(s, c + dx, cy + dy, DIA2, r=0.09 if i == 2 else 0.13)
        cross(s, c, cy, INK)
        caption(s, c - 0.95, cy + 0.58, 1.9,
                "same centre, still turned" if i == 1 else "orange sits inside blue",
                MUTED if i == 1 else GOOD, 11, True)
    caption(s, px - 0.2, 4.45, 3.9, sub.split("\n")[0], MUTED, 11.5)
    if "\n" in sub:
        caption(s, px - 0.2, 4.68, 3.9, sub.split("\n")[1], MUTED, 11.5)

arrow(s, 4.35, 3.65, 4.85, 3.65, MUTED, 2.0)
arrow(s, 8.45, 3.65, 8.95, 3.65, MUTED, 2.0)

caption(s, 0.7, 5.02, 12.0,
        "Subtracting the centre does not measure the turn. It removes position, so that the turn "
        "is the only thing left to find.",
        INK, 14, True, PP_ALIGN.LEFT)
caption(s, 0.7, 5.36, 12.0,
        "The markers are bolted to the pelvis, so they never move relative to each other, "
        "and their average is a fixed point on the body.",
        MUTED, 13, False, PP_ALIGN.LEFT)
caption(s, 0.7, 5.58, 12.0,
        "Subtract it from every marker and the robot could be anywhere in the room, the numbers "
        "would be identical. Only turning changes them.",
        MUTED, 13, False, PP_ALIGN.LEFT)
box(s, 0.7, 6.18, 12.0, 0.62,
    "All of this holds only while all four markers are visible",
    fill=PALE, line=ACC, color=ACC, size=15, bold=True)


# ================================================================== 3c. what went wrong
s = prs.slides.add_slide(BLANK)
header(s, "bug 1 explained  \u00b7  part 2", "What went wrong")

# one picture: a marker drops, the centre moves
BIG = [(-1.05, -0.75), (1.05, -0.75), (-1.05, 0.75), (1.05, 0.75)]
cx2, cy2 = 2.85, 3.55
for k, (dx, dy) in enumerate(BIG):
    dot(s, cx2 + dx, cy2 + dy, GREY if k == 2 else ACC, r=0.16, fill=WHITE)
caption(s, cx2 - 1.78, cy2 + 0.6, 1.5, "dropped", GREY, 11.5, True)
cross(s, cx2, cy2, GREY, r=0.12)
cross(s, cx2 + 0.35, cy2 - 0.25, BAD, r=0.12)
arrow(s, cx2, cy2, cx2 + 0.35, cy2 - 0.25, BAD, 2.25)
caption(s, cx2 + 0.5, cy2 - 0.62, 2.6, "centre moves 29 mm", BAD, 12.5, True, PP_ALIGN.LEFT)
caption(s, 0.7, 5.3, 5.9, "Lose a marker and the centre of what is left", MUTED, 13.5, False, PP_ALIGN.LEFT)
caption(s, 0.7, 5.58, 5.9, "is no longer the same point on the body.", MUTED, 13.5, False, PP_ALIGN.LEFT)

box(s, 7.0, 2.4, 5.7, 0.85,
    "The two shapes ended up measured\nabout centres 29 mm apart",
    fill=PALEB, line=BAD, color=BAD, size=15, bold=True)
box(s, 7.0, 3.65, 5.7, 0.85,
    "Fix: re-centre the recorded shape\non the markers that remain",
    fill=PALEG, line=GOOD, color=GOOD, size=15, bold=True)
box(s, 7.0, 4.9, 5.7, 0.95,
    "predicted 29.08 mm    measured 29.08 mm\nafter the fix:  0.38 mm",
    fill=WHITE, line=GREY, color=INK, size=14, bold=True)

caption(s, 0.7, 6.5, 12.0,
        "With all four markers visible nothing changes, so 67,427 full frames came out bit-identical.",
        MUTED, 13, False, PP_ALIGN.LEFT)


# ================================================================== 3d. bug 2
s = prs.slides.add_slide(BLANK)
header(s, "bug 2 explained", "A frame rate that was quietly rounded down")

box(s, 0.7, 1.75, 3.5, 0.72, "we record at  48.88 Hz", fill=PALE, line=ACC, color=INK, size=15, bold=True)
arrow(s, 4.35, 2.11, 5.05, 2.11, MUTED, 2.0)
box(s, 5.2, 1.75, 3.0, 0.72, "ASAP rounds it DOWN", fill=WHITE, line=GREY, color=MUTED, size=14)
arrow(s, 8.35, 2.11, 9.05, 2.11, BAD, 2.0)
box(s, 9.2, 1.75, 3.5, 0.72, "played as  48 Hz", fill=PALEB, line=BAD, color=BAD, size=15, bold=True)

# the two clocks pulling apart
ty, ty2 = 3.15, 3.75
for k in range(13):
    box(s, 0.95 + k * 0.40, ty, 0.055, 0.32, "", fill=ACC, line=ACC, shape=MSO_SHAPE.RECTANGLE)
    box(s, 0.95 + k * 0.425, ty2, 0.055, 0.32, "", fill=BAD, line=BAD, shape=MSO_SHAPE.RECTANGLE)
caption(s, 6.1, ty - 0.05, 1.0, "real", ACC, 12, True, PP_ALIGN.LEFT)
caption(s, 6.1, ty2 - 0.05, 1.3, "ASAP", BAD, 12, True, PP_ALIGN.LEFT)
caption(s, 0.7, 4.28, 6.2, "Over one whole dab, 298 frames or 6.1 seconds:", INK, 13.5, True, PP_ALIGN.LEFT)
box(s, 0.7, 4.62, 2.95, 0.62, "stored as 48\n5.4 frames behind",
    fill=PALEB, line=BAD, color=BAD, size=13.5, bold=True)
box(s, 3.75, 4.62, 2.95, 0.62, "stored as 49\n0.7 frames off",
    fill=PALEG, line=GOOD, color=GOOD, size=13.5, bold=True)

caption(s, 7.4, 2.85, 5.3, "Caught by asking for frames back", INK, 14, True, PP_ALIGN.LEFT)
caption(s, 7.4, 4.95, 5.3, "Ask for frame 250, get 245. The drift, mid-dab.", MUTED, 12.5, False, PP_ALIGN.LEFT)
table(s, [
    ("Asked for", "Got back"),
    ("10", "10"),
    ("100", ("98", BAD, True)),
    ("250", ("245", BAD, True)),
], top=3.25, left=7.4, col_w=[Inches(2.6), Inches(2.7)], size=14)

box(s, 0.7, 5.35, 12.0, 0.8,
    "Fix: store the rate as a whole number, 49",
    fill=PALEG, line=GOOD, color=GOOD, size=16, bold=True)
caption(s, 0.7, 6.35, 12.0,
        "A wrong clock never shows up in a training curve. Only asking for the data back reveals it.",
        MUTED, 13.5, False, PP_ALIGN.LEFT)


# ================================================================== 3e. bug 3
s = prs.slides.add_slide(BLANK)
header(s, "bug 3 explained", "Measuring from the wrong point on the body")

# the two points, one picture
box(s, 1.3, 2.3, 3.6, 2.1, "", fill=PALE, line=ACC)
caption(s, 1.3, 2.38, 3.6, "PELVIS", ACC, 12.5, True)
ox, oy = 3.1, 2.95
mx, my = 3.1, 3.85
for dx, dy in [(-0.7, -0.3), (0.7, -0.3), (-0.7, 0.3), (0.7, 0.3)]:
    dot(s, mx + dx, my + dy, DIA2, r=0.13)
cross(s, ox, oy, ACC, r=0.14)
cross(s, mx, my, DIA2, r=0.14)
arrow(s, ox, oy + 0.17, mx, my - 0.17, BAD, 2.25)
caption(s, ox + 0.16, oy + 0.22, 1.6, "83 mm", BAD, 14, True, PP_ALIGN.LEFT)
caption(s, 1.3, 4.55, 3.6, "what the simulator wants  \u25cf  above", ACC, 11.5, True)
caption(s, 1.3, 4.8, 3.6, "what the cameras report  \u25cf  below", DIA2, 11.5, True)

caption(s, 6.6, 2.32, 6.1, "The same height, measured two ways", INK, 14, True, PP_ALIGN.LEFT)
caption(s, 6.6, 2.66, 6.1, "How high is the pelvis above the floor?", MUTED, 12.5, False, PP_ALIGN.LEFT)
table(s, [
    ("", "height"),
    ("From the robot's own joint angles", "0.767 m"),
    ("From the cameras", "0.684 m"),
    (("Difference", BAD, True), ("0.083 m", BAD, True)),
], top=3.08, left=6.6, col_w=[Inches(4.0), Inches(2.1)], size=14)
caption(s, 6.6, 4.72, 6.1, "Same distance, but measured to two different points",
        MUTED, 12.5, False, PP_ALIGN.LEFT)
caption(s, 6.6, 4.94, 6.1, "on the body. The gap between them is the offset.",
        MUTED, 12.5, False, PP_ALIGN.LEFT)
box(s, 6.6, 5.3, 6.1, 0.62,
    "Fix: shift by that 83 mm, turned with the body",
    fill=PALEG, line=GOOD, color=GOOD, size=14, bold=True)

caption(s, 0.7, 6.1, 12.0,
        "Both capture sessions gave the same answer to within 1 mm, which is what makes it a fixed "
        "offset rather than noise.", MUTED, 13, False, PP_ALIGN.LEFT)


# ================================================================== 3. DIAGRAM: measurement
s = prs.slides.add_slide(BLANK)
header(s, "how the numbers are made", "Measuring the gap: replay real actions in simulation")

box(s, 0.7, 2.75, 2.2, 1.0, "REAL ROBOT\n473 dab reps recorded", fill=PALEG, line=GOOD, size=13, bold=True)

box(s, 3.35, 1.85, 2.4, 0.78, "initial state s\u2080\nwhere the sim starts", fill=PALEG, line=GOOD, size=12.5)
box(s, 3.35, 2.86, 2.4, 0.78, "actions a\u209c\nwhat we commanded", fill=WHITE, size=12.5)
box(s, 3.35, 4.55, 2.4, 0.78, "states s\u209c\nwhat the robot did", fill=WHITE, size=12.5)
arrow(s, 2.9, 3.05, 3.3, 2.24)
arrow(s, 2.9, 3.25, 3.3, 3.25)
arrow(s, 2.9, 3.45, 3.3, 4.94)

box(s, 6.2, 2.2, 2.5, 1.1, "SIMULATOR\nstart at s\u2080, then apply a\u209c", fill=PALE, line=ACC, size=13, bold=True)
arrow(s, 5.75, 2.24, 6.15, 2.55)
arrow(s, 5.75, 3.25, 6.15, 2.95)

box(s, 9.3, 2.35, 2.3, 0.8, "sim states \u015d\u209c", fill=WHITE, size=13)
arrow(s, 8.7, 2.75, 9.25, 2.75)

box(s, 9.3, 4.4, 3.3, 0.95, "| \u015d\u209c \u2212 s\u209c |  per joint\n= THE GAP",
    fill=PALEB, line=BAD, color=BAD, size=13, bold=True)
arrow(s, 10.45, 3.2, 10.45, 4.35, BAD)
arrow(s, 5.8, 4.94, 9.25, 4.94, BAD)

caption(s, 6.0, 3.42, 2.9, "same start, same actions", MUTED, 11.5, True)

bullets(s, [
    ("The sim is reset to the robot's real state first, so both start from the same place.", INK, True),
    "Then the same actions go into both. Any difference after that is the simulator being wrong.",
    "Re-anchored every 1.0 s, so small errors cannot pile up and look like a physics gap.",
], top=5.65, size=14, gap=5, height=1.5)


# ================================================================== 4. gap result
s = prs.slides.add_slide(BLANK)
header(s, "result 1", "The gap is the ankles, and almost nothing else")
caption(s, 0.7, 1.78, 7.0, "Error as a share of each joint's own range of motion",
        MUTED, 12.5, False, PP_ALIGN.LEFT)
data = [("L_AnklePitch", 43.2, True), ("L_AnkleRoll", 25.0, True), ("R_AnklePitch", 21.3, True),
        ("R_HipRoll", 20.6, False), ("R_AnkleRoll", 15.4, True), ("L_HipYaw", 10.9, False),
        ("L_HipPitch", 7.1, False), ("L_Knee", 3.5, False), ("L_ShoulderPitch", 2.8, False)]
mx = 43.2; x0 = 2.75; wmax = 4.3
for i, (name, val, is_ankle) in enumerate(data):
    yy = 2.25 + i * 0.42
    caption(s, 0.5, yy - 0.02, 2.15, name, INK if is_ankle else MUTED, 12, is_ankle, PP_ALIGN.RIGHT)
    w = max(val / mx * wmax, 0.06)
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x0), Inches(yy), Inches(w), Inches(0.26))
    sh.fill.solid(); sh.fill.fore_color.rgb = BAD if is_ankle else GREY
    sh.line.fill.background(); sh.shadow.inherit = False
    caption(s, x0 + w + 0.08, yy - 0.04, 1.0, "%.1f%%" % val,
            INK if is_ankle else MUTED, 11.5, is_ankle, PP_ALIGN.LEFT)
caption(s, 0.5, 6.12, 7.6, "All four ankle joints land in the top six.  On average they are 4.79x "
        "worse than every other joint.", INK, 13, True, PP_ALIGN.LEFT)

# ---- why measure it this way: two joints, same bar length
caption(s, 8.7, 1.78, 4.1, "Why measure it this way", INK, 13.5, True, PP_ALIGN.LEFT)
caption(s, 8.7, 2.1, 4.1, "Each bar is one joint's full travel.", MUTED, 11.5, False, PP_ALIGN.LEFT)
caption(s, 8.7, 2.32, 4.1, "Red is how wrong the simulator is.", MUTED, 11.5, False, PP_ALIGN.LEFT)

BW = 3.85
for (lab, moves, err, pct, y) in [("L_AnklePitch", 25.4, 11.0, 43.2, 2.95),
                                  ("L_ShoulderPitch", 135.3, 3.8, 2.8, 4.35)]:
    caption(s, 8.7, y, 4.1, "%s  moves %.0f\u00b0" % (lab, moves), INK, 12, True, PP_ALIGN.LEFT)
    base = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.7), Inches(y + 0.32),
                              Inches(BW), Inches(0.34))
    base.fill.solid(); base.fill.fore_color.rgb = PALE
    base.line.color.rgb = GREY; base.line.width = Pt(0.75); base.shadow.inherit = False
    ov = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.7), Inches(y + 0.32),
                            Inches(max(BW * pct / 100, 0.07)), Inches(0.34))
    ov.fill.solid(); ov.fill.fore_color.rgb = BAD
    ov.line.fill.background(); ov.shadow.inherit = False
    caption(s, 8.7, y + 0.74, 4.1, "off by %.1f\u00b0  of  %.0f\u00b0   =   %.0f%%" % (err, moves, pct),
            BAD if pct > 10 else MUTED, 12, True, PP_ALIGN.LEFT)

box(s, 8.7, 5.62, 3.85, 0.78,
    "Same size error in degrees\nmeans very different things",
    fill=PALE, line=ACC, color=ACC, size=12.5, bold=True)

note(s, "The paper asserts the G1's parallel-linkage ankle is the problem. This measures it, on our "
        "robot, before fitting anything.", top=6.72, size=13)


# ================================================================== 5. DIAGRAM: why 4 beats 23
s = prs.slides.add_slide(BLANK)
header(s, "the key idea", "Why correcting 4 joints beats correcting 23")
caption(s, 0.7, 1.72, 12.0, "The delta model is scored mostly on WHERE THE BODY IS "
        "(weight 3.0), barely on joint angles (0.5). That leaves a shortcut.",
        INK, 14.5, True, PP_ALIGN.LEFT)

# left panel: 23 DoF
box(s, 0.7, 2.35, 5.75, 0.5, "23 DoF free to correct", fill=PALEB, line=BAD, color=BAD, size=14, bold=True)
for i, (jn, free) in enumerate([("Hip", True), ("Knee", True), ("Ankle", True)]):
    box(s, 1.0 + i * 1.85, 3.05, 1.6, 0.62, jn,
        fill=PALEB if free else PALE, line=BAD if free else GREY,
        color=BAD if free else MUTED, size=13, bold=free)
caption(s, 0.7, 3.78, 5.75, "all free  →  many ways to place the pelvis", MUTED, 12)
box(s, 0.7, 4.2, 5.75, 0.62, "Bends the KNEE to put the body in the right place",
    fill=WHITE, line=BAD, color=INK, size=13)
box(s, 0.7, 4.95, 2.75, 0.62, "Body position  ✓", fill=PALEG, line=GOOD, color=GOOD, size=13, bold=True)
box(s, 3.7, 4.95, 2.75, 0.62, "Ankle still wrong  ✗", fill=PALEB, line=BAD, color=BAD, size=13, bold=True)

# right panel: 4 DoF
box(s, 6.9, 2.35, 5.75, 0.5, "4 DoF free (ankles only)", fill=PALEG, line=GOOD, color=GOOD, size=14, bold=True)
for i, (jn, free) in enumerate([("Hip", False), ("Knee", False), ("Ankle", True)]):
    box(s, 7.2 + i * 1.85, 3.05, 1.6, 0.62, jn,
        fill=PALEG if free else PALE, line=GOOD if free else GREY,
        color=GOOD if free else GREY, size=13, bold=free)
caption(s, 6.9, 3.78, 5.75, "locked  →  only one way to place the pelvis", MUTED, 12)
box(s, 6.9, 4.2, 5.75, 0.62, "Must actually fix the ANKLE to place the body",
    fill=WHITE, line=GOOD, color=INK, size=13)
box(s, 6.9, 4.95, 2.75, 0.62, "Body position  ✓", fill=PALEG, line=GOOD, color=GOOD, size=13, bold=True)
box(s, 9.9, 4.95, 2.75, 0.62, "Ankle fixed  ✓", fill=PALEG, line=GOOD, color=GOOD, size=13, bold=True)

note(s, "The leg is kinematically redundant: many hip/knee/ankle combinations put the pelvis in the same place. "
        "Constraining the model removes the shortcut and forces it to model the real physics.", top=5.85, size=14)

# ================================================================== 6. numbers
s = prs.slides.add_slide(BLANK)
header(s, "result 2", "4 DoF wins. 23 DoF is worse than no correction at all.")
table(s, [
    ("Delta model", "Joint error @1 s", "Knee error", "Verdict"),
    ("No correction (baseline)", "2.49 deg", "1.00 deg", ""),
    (("4-DoF, ankles only", GOOD, True), ("1.46 deg", GOOD, True), ("0.99 deg", GOOD, True), ("best", GOOD, True)),
    ("23-DoF, seed 1", ("3.30 deg", BAD, True), ("7.95 deg", BAD, True), ("worse than nothing", BAD, True)),
    ("23-DoF, seed 2", ("3.10 deg", BAD, True), ("3.74 deg", BAD, True), ("worse than nothing", BAD, True)),
], top=2.15, col_w=[Inches(3.9), Inches(2.7), Inches(2.4), Inches(3.0)], size=14.5)
bullets(s, [
    ("Look at the knees. That is the cheating, visible as a number.", INK, True),
    "The 4-DoF model leaves them at 0.99 deg against a 1.00 deg baseline: untouched, because it cannot touch them.",
    "Both 23-DoF seeds wreck them (7.95 and 3.74) while fixing the ankles only half as well.",
    "",
    ("The paper picks 4 DoF because it lacked data. We have the data it says is needed, and 4 DoF still wins.", ACC, True),
], top=4.3, size=15, gap=5, height=2.2)

# ================================================================== 7. reward trap
s = prs.slides.add_slide(BLANK)
header(s, "a trap worth knowing", "Training reward picks the wrong model")
table(s, [
    ("Run", "Training reward", "Error vs the real robot", "Verdict"),
    (("23-DoF seed 2", INK, True), ("5.34  ← highest", GOOD, True), ("3.10°", BAD, True), ("2nd worst", BAD, True)),
    ("23-DoF seed 1", "4.82", ("3.30°", BAD, True), ("worst", BAD, True)),
    (("no delta at all", MUTED, False), ("—", MUTED, False), ("2.49°", MUTED, True), ("", MUTED, False)),
    (("4-DoF ankles", INK, True), "4.78  ← lowest", ("1.46°", GOOD, True), ("best", GOOD, True)),
], top=2.25, col_w=[Inches(3.3), Inches(3.1), Inches(3.2), Inches(2.4)], size=14.5)
bullets(s, [
    ("Reward says seed 2 is the best run. Replay against the real robot says it is nearly the worst.", BAD, True),
    "Both 23-DoF models are worse than applying no correction at all.",
    "Choosing a checkpoint by reward, the normal thing to do, picks the wrong one.",
    "",
    ("Reward is measured in simulation. Quality can only be measured against the real robot.", ACC, True),
], top=4.7, size=15, gap=6, height=1.9)

# ================================================================== 8. DIAGRAM: deploy failure
s = prs.slides.add_slide(BLANK)
header(s, "result 3  ·  the problem", "Why the fine-tuned policy falls over")
caption(s, 0.7, 1.68, 12.0, "The delta model is deleted at deployment. The policy was never trained without it.",
        INK, 14.5, True, PP_ALIGN.LEFT)
# training row
caption(s, 0.7, 2.2, 2.2, "IN TRAINING", ACC, 12.5, True, PP_ALIGN.LEFT)
box(s, 0.7, 2.6, 2.3, 0.75, "POLICY\nankle 0.51", fill=PALE, line=ACC, size=13, bold=True)
box(s, 3.5, 2.6, 2.3, 0.75, "DELTA MODEL\nankle 1.07", fill=PALEB, line=BAD, color=BAD, size=13, bold=True)
box(s, 6.4, 2.6, 2.5, 0.75, "TOTAL 1.58\ninto the ankles", fill=PALEG, line=GOOD, color=GOOD, size=13, bold=True)
box(s, 9.5, 2.6, 3.2, 0.75, "policy learns to balance\nin THIS world", fill=WHITE, line=GREY, size=12.5)
arrow(s, 3.0, 2.97, 3.45, 2.97); arrow(s, 5.85, 2.97, 6.35, 2.97); arrow(s, 8.95, 2.97, 9.45, 2.97)
# deploy row
caption(s, 0.7, 3.85, 2.2, "AT DEPLOYMENT", BAD, 12.5, True, PP_ALIGN.LEFT)
box(s, 0.7, 4.25, 2.3, 0.75, "POLICY\nankle 0.51", fill=PALE, line=ACC, size=13, bold=True)
box(s, 3.5, 4.25, 2.3, 0.75, "DELETED", fill=WHITE, line=GREY, color=GREY, size=13, bold=True)
box(s, 6.4, 4.25, 2.5, 0.75, "TOTAL 0.51\n32% of training", fill=PALEB, line=BAD, color=BAD, size=13, bold=True)
box(s, 9.5, 4.25, 3.2, 0.75, "falls over", fill=PALEB, line=BAD, color=BAD, size=13, bold=True)
arrow(s, 3.0, 4.62, 3.45, 4.62, GREY); arrow(s, 5.85, 4.62, 6.35, 4.62, BAD)
arrow(s, 8.95, 4.62, 9.45, 4.62, BAD)
bullets(s, [
    ("The delta was supposed to be a small correction. It became a co-controller.", BAD, True),
    "Its penalty term, exp(−||Δa||), goes flat once the correction is large, so nothing pulls it back down.",
    "The correction grew from 3.7 to 10.7 during training, with the penalty reading exactly zero.",
], top=5.45, size=15, gap=5, height=1.6)

# ================================================================== 9. status
s = prs.slides.add_slide(BLANK)
header(s, "where things stand", "Solid, and not solid")
box(s, 0.7, 1.95, 5.9, 0.5, "SOLID", fill=PALEG, line=GOOD, color=GOOD, size=14, bold=True)
bullets(s, [
    "Capture pipeline fixed and validated",
    "   473 reps, 140,821 frames, fully reproducible",
    "Sim-to-real gap measured",
    "   ankles 4.79x, reproduces to ~1%",
    "4-DoF beats 23-DoF",
    "   two seeds, independent evaluation samples",
], top=2.6, left=0.9, width=5.7, size=14, gap=3, height=3.4)
box(s, 7.0, 1.95, 5.7, 0.5, "NOT SOLID", fill=PALEB, line=BAD, color=BAD, size=14, bold=True)
bullets(s, [
    "No deployable policy yet",
    "   both fine-tuned policies fall over in MuJoCo",
    "Simulation cannot pick between them",
    "   each policy wins in the simulator it trained in",
    "Only one motion tested",
    "   a standing, arm-heavy dab",
], top=2.6, left=7.2, width=5.4, size=14, gap=3, height=3.4)
note(s, "v6 remains the only policy validated on hardware. Nothing on the robot has changed this week.",
     top=5.6, color=INK, size=15)

# ================================================================== 10. next
s = prs.slides.add_slide(BLANK)
header(s, "next", "Three things, in order")
for i, (n, t1, t2, col) in enumerate([
    ("1", "Shrink the delta model",
     "Replace exp(−||Δa||) with a penalty that does not go flat, so the correction stays a correction.\n"
     "Nothing else works until the delta is small enough to remove safely.", BAD),
    ("2", "Re-fine-tune against the smaller delta",
     "Then check it in MuJoCo. That is the gate. Nothing goes on the robot until it passes.", ACC),
    ("3", "Hardware A/B against v6",
     "The only test that answers whether any of this transfers. Needs robot time.", ACC),
]):
    y = 2.05 + i * 1.5
    box(s, 0.7, y, 0.55, 0.75, n, fill=WHITE, line=col, color=col, size=17, bold=True)
    tf = tb(s, 1.45, y - 0.05, 11.3, 0.45)
    run(tf.paragraphs[0], t1, 19, True, INK, after=2)
    tf2 = tb(s, 1.45, y + 0.38, 11.3, 0.85)
    for j, ln in enumerate(t2.split("\n")):
        p = tf2.paragraphs[0] if j == 0 else tf2.add_paragraph()
        run(p, ln, 14, False, MUTED, after=2)
note(s, "Domain randomisation was tried this week and made the dependence worse, not better. It is not the fix.",
     top=6.6, color=MUTED, size=13.5)

out = os.path.join(HERE, "asap_week_20260809.pptx")
prs.save(out)
print(f"wrote {out}")
