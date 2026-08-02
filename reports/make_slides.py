#!/usr/bin/env python3
"""Generate the ASAP progress deck as .pptx (Drive converts it to Google Slides).

  /usr/bin/python3 make_slides.py
Writes asap_progress_20260730.pptx next to this script.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))

INK   = RGBColor(0x1A, 0x1D, 0x24)
MUTED = RGBColor(0x5B, 0x63, 0x72)
ACC   = RGBColor(0x1F, 0x6F, 0x8B)
GOOD  = RGBColor(0x1B, 0x7F, 0x53)
BAD   = RGBColor(0xA6, 0x3D, 0x40)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def tb(slide, l, t, w, h):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def title_slide(title, subtitle, kicker=None):
    s = prs.slides.add_slide(BLANK)
    if kicker:
        tf = tb(s, 0.9, 1.7, 11.5, 0.4)
        p = tf.paragraphs[0]; r = p.add_run(); r.text = kicker.upper()
        r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = ACC
    tf = tb(s, 0.9, 2.2, 11.5, 1.6)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = INK
    tf2 = tb(s, 0.9, 3.9, 11.0, 1.2)
    p = tf2.paragraphs[0]; r = p.add_run(); r.text = subtitle
    r.font.size = Pt(18); r.font.color.rgb = MUTED
    return s


def section(title, num):
    s = prs.slides.add_slide(BLANK)
    tf = tb(s, 0.9, 2.9, 11.5, 0.5)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = num
    r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = ACC
    tf2 = tb(s, 0.9, 3.35, 11.5, 1.0)
    p = tf2.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = INK
    return s


def content(title, bullets, note=None):
    s = prs.slides.add_slide(BLANK)
    tf = tb(s, 0.9, 0.65, 11.5, 0.9)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = INK

    body = tb(s, 0.9, 1.75, 11.5, 4.7)
    first = True
    for b in bullets:
        # accepts "text", (level, "text"), ("text", level), each with optional style
        if isinstance(b, str):
            lvl, text, style = 0, b, ""
        else:
            parts = list(b)
            if isinstance(parts[0], int):
                lvl, text = parts[0], parts[1]
            else:
                text = parts[0]
                lvl = parts[1] if len(parts) > 1 and isinstance(parts[1], int) else 0
            style = next((p for p in parts if isinstance(p, str)
                          and p in ("good", "bad", "muted")), "")
        p = body.paragraphs[0] if first else body.add_paragraph()
        first = False
        p.level = lvl
        r = p.add_run()
        r.text = ("• " if lvl == 0 else "– ") + text
        r.font.size = Pt(17 if lvl == 0 else 15)
        r.font.color.rgb = {"good": GOOD, "bad": BAD, "muted": MUTED}.get(style, INK)
        if style == "good":
            r.font.bold = True
        p.space_after = Pt(10 if lvl == 0 else 6)
    if note:
        nf = tb(s, 0.9, 6.5, 11.5, 0.6)
        p = nf.paragraphs[0]; r = p.add_run(); r.text = note
        r.font.size = Pt(13); r.font.italic = True; r.font.color.rgb = MUTED
    return s


def metrics(title, rows, note=None, headers=("Before", "After")):
    s = prs.slides.add_slide(BLANK)
    tf = tb(s, 0.9, 0.65, 11.5, 0.9)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = INK

    shape = s.shapes.add_table(len(rows) + 1, 3, Inches(0.9), Inches(1.9),
                               Inches(11.5), Inches(0.5 * (len(rows) + 1)))
    t = shape.table
    t.columns[0].width = Inches(5.1)
    t.columns[1].width = Inches(3.2)
    t.columns[2].width = Inches(3.2)
    for j, h in enumerate(["", headers[0], headers[1]]):
        c = t.cell(0, j); c.text = h or " "
        pr = c.text_frame.paragraphs[0]
        if pr.runs:
            pr.runs[0].font.size = Pt(14); pr.runs[0].font.bold = True
            pr.runs[0].font.color.rgb = MUTED
        if j:
            pr.alignment = PP_ALIGN.CENTER
    for i, (lab, before, after) in enumerate(rows, 1):
        for j, val in enumerate([lab, before, after]):
            c = t.cell(i, j); c.text = val or " "
            pr = c.text_frame.paragraphs[0]
            if not pr.runs:
                continue
            run = pr.runs[0]
            run.font.size = Pt(15)
            if j == 0:
                run.font.color.rgb = INK
            else:
                pr.alignment = PP_ALIGN.CENTER
                run.font.bold = True
                run.font.color.rgb = BAD if j == 1 else GOOD
    if note:
        nf = tb(s, 0.9, 1.95 + 0.5 * (len(rows) + 1), 11.5, 0.7)
        p = nf.paragraphs[0]; r = p.add_run(); r.text = note
        r.font.size = Pt(13); r.font.italic = True; r.font.color.rgb = MUTED
    return s


def picture(title, img, note=None):
    s = prs.slides.add_slide(BLANK)
    tf = tb(s, 0.9, 0.5, 11.5, 0.8)
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = INK
    if os.path.exists(img):
        s.shapes.add_picture(img, Inches(1.1), Inches(1.45), width=Inches(11.1))
    if note:
        nf = tb(s, 0.9, 6.75, 11.5, 0.6)
        p = nf.paragraphs[0]; r = p.add_run(); r.text = note
        r.font.size = Pt(12); r.font.italic = True; r.font.color.rgb = MUTED
    return s


# ----------------------------------------------------------------- the deck
title_slide(
    "Real-world data collection for ASAP on the Unitree G1",
    "Capture pipeline built, first clean dataset collected, calibrated and validated.\nEric Xie  ·  30 July 2026",
    kicker="Progress update")

content("Why we need this", [
    "ASAP closes the sim-to-real gap by learning a delta-action model from REAL rollouts",
    "Training that model needs, at every control step:",
    (1, "joint state and the action the policy actually commanded"),
    (1, "the floating base (pelvis) pose — which the robot cannot measure itself"),
    "The onboard IMU gives orientation but NOT position, so an external reference is required",
    "Chosen source: Vicon motion capture of a 4-marker cluster on the pelvis",
], note="The hard part is not measuring either stream — it is putting them on one trustworthy timeline.")

section("What was built", "01")

content("Capture pipeline", [
    "Vicon Nexus streams LABELLED pelvis markers over OSC at 100 Hz",
    "The robot's Jetson receives them and timestamps on arrival",
    ("rl_sar records state + action at each policy step (~50 Hz) on the SAME machine", 0, "good"),
    "Pelvis pose computed from the markers by a least-squares Umeyama/Kabsch rigid fit",
    "Streams merged offline by timestamp; pose interpolated to each step's exact instant",
], note="Tooling: scripts/vicon/ — osc_listener.py, merge_rollout_vicon.py, clean_dataset.py")

content("The key design decision: one clock", [
    ("Both streams are timestamped by the same machine, so only ONE clock exists", 0, "good"),
    "This removes an entire class of problem:",
    (1, "no NTP / PTP clock synchronisation"),
    (1, "no clock-offset or drift estimation"),
    (1, "no sync event (clap, stomp, shove) to cross-correlate"),
    "Alignment is plain timestamp matching — pose interpolated to the step's own instant",
    "Unlike a latest-value buffer, dropouts stay visible rather than being silently filled",
    (1, "with a stale pose, so occluded spans can be excluded instead of trusted"),
])

section("Problems found and fixed", "02")

content("1 — Right ankle wind-up destroyed early runs", [
    ("Symptom: robot cut power mid-dab, repeatedly (firmware overload protection)", 0, "bad"),
    "Traced in the logs to R_ankle_roll being commanded to ~100° against a ~15° joint limit",
    "Root cause: ankle_roll is an under-constrained action dimension",
    (1, "planted feet mean it barely affects the tracking reward, so training let it drift"),
    (1, "simulation hides it — joint limits, contact and the force cap absorb the command"),
    (1, "the reference motion asks for 0° ankle roll on both feet, for all 179 frames"),
    ("Fixed by a targeted fine-tune penalising commanded targets outside joint limits", 0, "good"),
], note="A deployment-side action clamp kept collection going until the fine-tune landed.")

metrics("1 — Result of the ankle fine-tune", [
    ("R_ankle_roll commanded", "~100°", "1.7–2.9°"),
    ("R_ankle_roll torque", "~30 Nm", "0.6–3.5 Nm"),
    ("Power cuts per session", "every run", "none"),
    ("Consecutive dab reps", "0–2, then dead", "30+"),
], note="Diagnosed from the captured logs and handed to the training side, which produced the fix.")

content("2 — Motion capture kept mislabelling the markers", [
    ("Symptom: the pelvis rigid body flipped ~180°, corrupting orientation", 0, "bad"),
    "The G1's silver shell throws IR reflections, spawning phantom markers",
    "The arm occludes pelvis markers during the dab, leaving the solve under-constrained",
    "Left/right markers sit only ~75 mm apart (vs ~170 mm front/back) — easy to confuse",
    ("Fixed by moving to Nexus and redefining the subject", 0, "good"),
    (1, "Nexus streams LABELLED markers, so identity is fixed and cannot be swapped"),
    (1, "axes defined from the robot's perspective: origin right-back, X forward, Y left"),
], note="An automatic geometry-based repair was tried first and failed — a 180°-rotated labelling fits a near-symmetric cluster equally well.")

content("3 — Robot kept losing communication", [
    ("Symptom: rt/lowstate stopped intermittently; robot needed repeated power cycles", 0, "bad"),
    "Network and DDS discovery both looked healthy, which made it hard to see",
    "Root cause: a second IP subnet had been added to the robot's network interface",
    (1, "the Jetson then advertised TWO DDS locators"),
    (1, "the robot intermittently sent data to the unreachable one"),
    ("Fixed by keeping one address on the interface and giving the Vicon PC an address", 0, "good"),
    (1, "on the robot subnet via a DHCP reservation — 6/6 stability windows afterwards"),
])

section("Result", "03")

metrics("Dataset quality — before vs after the fixes", [
    ("Orientation flips", "655", "0"),
    ("Mislabelled frames", "17.8%", "0.00%"),
    ("Markers tracked per frame", "dropouts", "4 / 4 always"),
    ("Rigid-body fit residual", "—", "0.40 mm median"),
    ("Vicon coverage of policy steps", "partial", "100%"),
], note="Same cameras, same robot — the difference is the subject definition and the network fix.")

content("The dataset", [
    ("69 complete dab repetitions — 20,556 transitions, 411 s of dab motion", 0, "good"),
    "Every step carries: action, joint position / velocity / torque, IMU, and pelvis pose",
    "100% mocap coverage — no step is missing a base pose",
    "Nothing dropped: none too short, none after a power cut, no coverage gaps",
    "Published as a release asset with a dataset card describing schema, quality and caveats",
], note="github.com/EricXieQ/rl_sar_g1_deployment — release capture-20260730")

content("Calibration to the robot's own frame", [
    "The Vicon marker frame and the URDF pelvis frame differ by a fixed rotation",
    "Solved by least squares against the IMU, which is in the URDF frame by definition",
    ("Result: roll +4.51°, pitch +5.99°, yaw −3.71°", 0, "good"),
    "Applying it brings Vicon and IMU orientation into agreement:",
    (1, "8.23° median disagreement → 2.63°"),
    "The small offset confirms the marker axes are correctly oriented; it also matches an",
    (1, "independent hand measurement taken days earlier — a genuine constant, not noise"),
])

section("Validation", "04")

content("How we know the data is right", [
    ("Vicon and the onboard IMU measure the same motion through completely separate", 0, "good"),
    (1, "physical paths — neither can influence the other, so agreement is real evidence"),
    "Angular velocity: Vicon-derived vs IMU gyroscope → correlation 0.92 (pitch), 0.94 (yaw)",
    "Gravity, measured two independent ways: +9.727 vs +9.719 m/s² — agree to 0.08%",
    ("Time alignment: cross-correlation peaks at exactly 0 ms — no residual offset", 0, "good"),
    "Repeatability: 69 reps of the same motion overlay to 1.4 mm median spread",
], note="Reproduce with reports/validate.py")

picture("Validation — independent sensor agreement",
        os.path.join(HERE, "validation.png"),
        note="Top: Vicon-derived angular velocity tracks the IMU gyroscope through fast transients. Bottom left: all 69 reps overlaid. Bottom right: time alignment peaks at zero lag.")

content("What this validation does NOT prove", [
    "Dynamic linear acceleration could not be validated directly",
    (1, "differentiating position twice at 50 Hz amplifies noise enormously —"),
    (1, "0.4 mm of position noise becomes ~1 m/s² of spurious acceleration"),
    (1, "so that comparison is noise-dominated (raw correlation 0.11–0.24)"),
    "Position is therefore validated indirectly — via gravity and via repeatability",
    "The roll axis correlates less well (0.60) simply because the pelvis barely rolls in a dab",
    (1, "little signal to correlate, not a defect in the measurement"),
], note="Stated explicitly so the delta-model results are read with the right error bars.")

section("Next", "05")

content("Next steps", [
    "Fit the delta-action model on the 69 repetitions and evaluate the corrected simulator",
    "Fine-tune the policy in the corrected sim and re-deploy",
    "Collect a second session with deliberate variety rather than more identical reps:",
    (1, "different positions and headings in the capture volume"),
    (1, "a few repetitions at reduced action scale"),
    (1, "small perturbations, so the model sees off-nominal states"),
    "Open question for the training side: should the reference dab end heading-neutral?",
    (1, "the robot yaws slightly per repetition — worth a symmetry or yaw-drift term"),
])

content("Known limitations", [
    "Vicon orientation needs the calibration applied; the IMU is already in the URDF frame",
    "IMU yaw drifts slowly over a long session — roll and pitch are gravity-referenced and do not",
    "Base velocities are finite-differenced from interpolated pose, so noisier than position",
    "All 69 repetitions are the same motion from broadly the same place — coverage, not count,",
    (1, "is what the next session should add"),
    "Only 9 of 20 mocap cameras are currently working, limiting the usable capture volume",
])

out = os.path.join(HERE, "asap_progress_20260730.pptx")
prs.save(out)
print("wrote %s  (%d slides)" % (out, len(prs.slides._sldIdLst)))
